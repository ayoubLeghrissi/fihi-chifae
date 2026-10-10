"""Sales analytics and Excel reporting.

Only *confirmed* orders count, attributed to the month they were confirmed.
Revenue = items + delivery fees actually charged.
"""

import datetime
import io

from . import catalog
from . import db as database

_CONFIRMED_ORDERS = (
    "SELECT o.id, substr(o.confirmed_at,1,7) AS ym, COALESCE(o.delivery_fee,0) AS fee,"
    " COALESCE((SELECT SUM(price*qty) FROM order_items WHERE order_id=o.id),0) AS items,"
    " COALESCE((SELECT SUM(qty) FROM order_items WHERE order_id=o.id),0) AS units"
    " FROM orders o WHERE o.status='confirmed'"
)


def _months_back(ym, count):
    year, month = (int(x) for x in ym.split("-"))
    out = []
    for _ in range(count):
        out.append(f"{year:04d}-{month:02d}")
        month -= 1
        if month == 0:
            year, month = year - 1, 12
    return list(reversed(out))


def _agg(rows):
    return {
        "orders": len(rows),
        "units": int(sum(r["units"] for r in rows)),
        "revenue": round(sum(r["items"] + r["fee"] for r in rows), 2),
        "delivery": round(sum(r["fee"] for r in rows), 2),
    }


def sales_summary():
    m = database.current_month()
    with database.live_db() as conn:
        confirmed = [dict(r) for r in conn.execute(_CONFIRMED_ORDERS)]
        pending = dict(
            conn.execute(
                "SELECT COUNT(*) AS orders,"
                " COALESCE(SUM((SELECT SUM(price*qty) FROM order_items WHERE order_id=o.id)"
                " + COALESCE(o.delivery_fee,0)),0) AS amount"
                " FROM orders o WHERE o.status='pending'"
            ).fetchone()
        )
        top = [
            dict(r)
            for r in conn.execute(
                "SELECT oi.product_name name, SUM(oi.qty) units,"
                " ROUND(SUM(oi.price*oi.qty),2) revenue FROM order_items oi"
                " JOIN orders o ON o.id = oi.order_id WHERE o.status='confirmed'"
                " GROUP BY oi.product_name ORDER BY revenue DESC LIMIT 6"
            )
        ]
        by_cat = [
            dict(r)
            for r in conn.execute(
                "SELECT oi.cat label, SUM(oi.qty) units,"
                " ROUND(SUM(oi.price*oi.qty),2) revenue FROM order_items oi"
                " JOIN orders o ON o.id = oi.order_id WHERE o.status='confirmed'"
                " GROUP BY oi.cat ORDER BY revenue DESC"
            )
        ]
        subscribers = conn.execute("SELECT COUNT(*) FROM newsletter").fetchone()[0]

    by_month = []
    for ym in _months_back(m, 6):
        by_month.append(dict(_agg([r for r in confirmed if r["ym"] == ym]), ym=ym))
    months_with_sales = sorted({r["ym"] for r in confirmed if r["ym"]}, reverse=True)

    return {
        "totals": _agg(confirmed),
        "month": dict(_agg([r for r in confirmed if r["ym"] == m]), label=m),
        "pending": {"orders": pending["orders"], "amount": round(pending["amount"] or 0, 2)},
        "revenueByMonth": by_month,
        "months": months_with_sales,
        "topSellers": top,
        "byCategory": by_cat,
        "categoryNames": catalog.category_names(),
        "subscribers": subscribers,
    }


def export_month(ym):
    sql = (
        "SELECT o.id order_id, o.confirmed_at, o.buyer, o.phone, o.address, o.city, o.location,"
        " COALESCE(o.delivery_fee,0) delivery_fee, o.source,"
        " oi.product_name, oi.cat, oi.size, oi.price, oi.qty"
        " FROM order_items oi JOIN orders o ON o.id = oi.order_id"
        " WHERE o.status='confirmed'"
    )
    args = ()
    if ym:
        sql += " AND substr(o.confirmed_at,1,7)=?"
        args = (ym,)
    sql += " ORDER BY o.confirmed_at, oi.id"
    with database.live_db() as conn:
        rows = [dict(r) for r in conn.execute(sql, args)]
        subscribers = [
            dict(r)
            for r in conn.execute("SELECT email,lang,created FROM newsletter ORDER BY created")
        ]
    return _build_excel(rows, subscribers, ym, catalog.category_names())


def _build_excel(rows, subscribers, ym, cat_names):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    head_font = Font(bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="16A34A")

    wb = Workbook()
    ws = wb.active
    ws.title = "المبيعات"
    ws.sheet_view.rightToLeft = True
    title = f"المبيعات المؤكدة - {ym}" if ym else "المبيعات المؤكدة - كل الأشهر"
    ws.cell(1, 1, title).font = Font(bold=True, size=13, color="FFFFFF")
    ws.cell(1, 1).fill = PatternFill("solid", fgColor="12843D")
    ws.cell(1, 2, f"تاريخ الإنشاء: {database.now()[:16].replace('T', ' ')}")

    headers = [
        "تاريخ التأكيد",
        "رقم الطلب",
        "المنتج",
        "الفئة",
        "الحجم",
        "سعر الوحدة (د.م.)",
        "الكمية",
        "المجموع (د.م.)",
        "الزبون",
        "الهاتف",
        "العنوان",
        "المدينة",
        "رابط الموقع",
        "المصدر",
    ]
    for col, h in enumerate(headers, start=1):
        c = ws.cell(2, col, h)
        c.font, c.fill = head_font, head_fill
        c.alignment = Alignment(horizontal="center")

    r = 3
    items_total = 0.0
    fees = {}
    for row in rows:
        total = (row["price"] or 0) * (row["qty"] or 0)
        items_total += total
        fees[row["order_id"]] = row["delivery_fee"] or 0
        values = [
            (row["confirmed_at"] or "")[:10],
            row["order_id"],
            row["product_name"],
            cat_names.get(row["cat"], row["cat"] or ""),
            row["size"] or "",
            round(row["price"] or 0, 2),
            row["qty"],
            round(total, 2),
            row["buyer"] or "",
            row["phone"] or "",
            row["address"] or "",
            row["city"] or "",
            row["location"] or "",
            "يدوي" if row["source"] == "manual" else "المتجر",
        ]
        for col, v in enumerate(values, start=1):
            ws.cell(r, col, v)
        r += 1

    fee_total = sum(fees.values())
    r += 1
    for label, value in (
        ("مجموع المنتجات", items_total),
        ("رسوم التوصيل", fee_total),
        ("الإجمالي", items_total + fee_total),
    ):
        ws.cell(r, 7, label).font = Font(bold=True)
        ws.cell(r, 8, round(value, 2)).font = Font(bold=True)
        r += 1

    widths = [14, 22, 34, 16, 10, 16, 9, 14, 24, 16, 30, 16, 30, 10]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    ns = wb.create_sheet("النشرة البريدية")
    ns.sheet_view.rightToLeft = True
    for col, h in enumerate(["البريد الإلكتروني", "اللغة", "تاريخ الاشتراك"], start=1):
        c = ns.cell(1, col, h)
        c.font, c.fill = head_font, head_fill
    for i, s in enumerate(subscribers, start=2):
        ns.cell(i, 1, s["email"])
        ns.cell(i, 2, s["lang"] or "")
        ns.cell(i, 3, (s["created"] or "")[:16].replace("T", " "))
    for i, w in enumerate([34, 8, 20], start=1):
        ns.column_dimensions[get_column_letter(i)].width = w

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    stamp = ym or datetime.date.today().strftime("all-%Y%m%d")
    return buf, f"ventes_{stamp}.xlsx"
